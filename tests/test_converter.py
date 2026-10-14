"""Tests for document converter (HTML, IPYNB, DOCX, PDF)."""

import json
from pathlib import Path

import docx
from pypdf import PdfWriter

from src.modules.converter import (
    convert_docx_to_markdown,
    convert_html_to_markdown,
    convert_ipynb_to_markdown,
    convert_pdf_to_markdown,
    convert_to_markdown,
)


def test_convert_html():
    raw_html = """
    <!DOCTYPE html>
    <html>
    <head><title>Test</title><script>alert('bad');</script></head>
    <body>
        <nav><a href="/">Home</a></nav>
        <h1>Main Heading</h1>
        <p>This is a paragraph with <b>bold</b> text.</p>
        <pre><code>def hello(): pass</code></pre>
        <ul><li>Item 1</li><li>Item 2</li></ul>
        <table>
            <tr><th>Name</th><th>Role</th></tr>
            <tr><td>Alice</td><td>Dev</td></tr>
        </table>
        <footer>Copyright 2026</footer>
    </body>
    </html>
    """
    md = convert_html_to_markdown(raw_html)
    assert "# Main Heading" in md
    assert "This is a paragraph" in md
    assert "alert('bad')" not in md
    assert "Copyright" not in md
    assert "| Name | Role |" in md
    assert "- Item 1" in md


def test_convert_ipynb(tmp_path: Path):
    nb_content = {
        "cells": [
            {
                "cell_type": "markdown",
                "source": ["## Data Science Step 1\n", "Analyze numbers."],
            },
            {
                "cell_type": "code",
                "source": ["import numpy as np\n", "print(np.sum([1, 2, 3]))"],
                "outputs": [{"text": ["6\n"]}],
            },
        ],
        "metadata": {},
        "nbformat": 4,
        "nbformat_minor": 2,
    }
    nb_path = tmp_path / "test.ipynb"
    nb_path.write_text(json.dumps(nb_content), encoding="utf-8")

    md = convert_ipynb_to_markdown(nb_path, include_outputs=False)
    assert "## Data Science Step 1" in md
    assert "import numpy as np" in md
    # Output excluded to save tokens
    assert "Output:" not in md

    md_with_out = convert_ipynb_to_markdown(nb_path, include_outputs=True)
    assert "Output:" in md_with_out
    assert "6" in md_with_out


def test_convert_docx(tmp_path: Path):
    doc = docx.Document()
    doc.add_heading("Docx Title", level=1)
    doc.add_paragraph("First paragraph text.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Header A"
    table.rows[0].cells[1].text = "Header B"
    table.rows[1].cells[0].text = "Val 1"
    table.rows[1].cells[1].text = "Val 2"

    docx_path = tmp_path / "test.docx"
    doc.save(str(docx_path))

    md = convert_docx_to_markdown(docx_path)
    assert "# Docx Title" in md
    assert "First paragraph text." in md
    assert "| Header A | Header B |" in md


def test_convert_pdf(tmp_path: Path):
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    pdf_path = tmp_path / "test.pdf"
    with open(pdf_path, "wb") as f:
        writer.write(f)

    md = convert_pdf_to_markdown(pdf_path)
    assert isinstance(md, str)


def test_convert_to_markdown_unified(tmp_path: Path):
    test_file = tmp_path / "notes.html"
    test_file.write_text("<h2>Notes</h2><p>Remember to test everything.</p>", encoding="utf-8")

    res = convert_to_markdown(test_file)
    assert res["file_name"] == "notes.html"
    assert "## Notes" in res["markdown"]
    assert res["reduced_tokens"] > 0


def test_convert_html_nested_blocks_not_duplicated():
    html = (
        "<ul><li><p>Hallo Welt</p><ul><li>Innen</li></ul></li></ul>"
        "<table><tr><td><p>Zelle</p></td></tr></table>"
    )
    md = convert_html_to_markdown(html)
    assert md.count("Hallo Welt") == 1
    assert md.count("Innen") == 1
    assert md.count("Zelle") == 1
